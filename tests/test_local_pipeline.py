"""
End-to-end tests of the legacy pipeline's local steps, against the
real dummy PDFs: extraction, report generation, and the DOCX schedule.

Needs two local-only folders that are deliberately not in the repo:
templates/ and "testing pdfs/". If either is missing, every test here
is skipped with a reason rather than failing, so the suite still
passes for anyone who hasn't been sent them.

Reports and schedules are written to pytest's temporary folders, not
reports/ or data/, so running this never touches real output.

Run from the project root:
    python -m pytest tests/test_local_pipeline.py -v
"""

from pathlib import Path

import pytest
from docx import Document

from app.pdf_extractor import extract_schedule_data
from app.schedule_generator import create_or_update_schedule_doc
from app.template_filler import fill_report_template


ROOT = Path(__file__).resolve().parents[1]
ADULT_TEMPLATE = ROOT / "templates" / "adult_report_template.docx"
CHILD_TEMPLATE = ROOT / "templates" / "child_report_template.docx"
PDF_DIR = ROOT / "testing pdfs"
ADULT_AGE_CUTOFF = 18

pytestmark = pytest.mark.skipif(
    not (ADULT_TEMPLATE.exists() and CHILD_TEMPLATE.exists() and PDF_DIR.is_dir()),
    reason="needs local templates/ and 'testing pdfs/' folders, which aren't in the repo",
)


def extract(pdf_name):
    path = PDF_DIR / pdf_name
    if not path.exists():
        pytest.skip(f"{pdf_name} not found in 'testing pdfs/'")
    return extract_schedule_data(str(path), ADULT_AGE_CUTOFF)


def document_text(doc):
    """All text in a docx's body paragraphs and tables."""
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def schedule_case_ids(schedule_path):
    table = Document(schedule_path).tables[0]
    return [row.cells[0].text for row in table.rows[1:]]


# --- Extraction -----------------------------------------------------


@pytest.mark.parametrize(
    "pdf_name, case_id, patient, dob, adult, child, schedule_date",
    [
        ("Test1-may4_2026.pdf", "15419897", "Anand Khanna", "November 4, 2011", "", "X", "May 4th, 2026"),
        ("Test2-may4_2026.pdf", "15419898", "Rajat Pawar", "November 4, 2001", "X", "", "May 4th, 2026"),
        ("Test3-may4_2026.pdf", "15419900", "Yash Dhamecha", "January 1, 1996", "X", "", "May 4th, 2026"),
        ("Test4-may10_2026.pdf", "15419372", "Rajat Pawar", "November 4, 2005", "X", "", "May 10th, 2026"),
    ],
)
def test_extraction_matches_known_values(pdf_name, case_id, patient, dob, adult, child, schedule_date):
    data = extract(pdf_name)
    assert data["case_id"] == case_id
    assert data["patient_name"] == patient
    assert data["dob"] == dob
    assert data["adult"] == adult
    assert data["child"] == child
    assert data["schedule_date"] == schedule_date
    assert data["total"] == "484.00"


def test_service_codes_extracted_whole():
    # The code includes the test abbreviation, e.g. "96101-WISC-V",
    # not just the five-digit billing code.
    data = extract("Test1-may4_2026.pdf")
    assert [s["code"] for s in data["services"]] == [
        "96101-WISC-V",
        "90791",
        "96130-AUTISM",
    ]


# --- Report generation ----------------------------------------------


@pytest.mark.parametrize(
    "pdf_name, template",
    [
        ("Test1-may4_2026.pdf", CHILD_TEMPLATE),
        ("Test2-may4_2026.pdf", ADULT_TEMPLATE),
    ],
)
def test_report_is_generated_with_no_placeholders_left(pdf_name, template, tmp_path):
    data = extract(pdf_name)
    report_path = Path(fill_report_template(data, str(template), str(tmp_path)))

    assert report_path.exists()
    assert report_path.parent == tmp_path

    text = document_text(Document(report_path))
    assert "{{" not in text, "a template placeholder was left unfilled"
    assert data["patient_name"].split()[0] in text


# --- DOCX schedule --------------------------------------------------


def test_schedule_rows_sorted_by_time_regardless_of_processing_order(tmp_path):
    # Process the 10 AM case before the 8 AM case.
    create_or_update_schedule_doc(extract("Test2-may4_2026.pdf"), str(tmp_path))
    schedule = create_or_update_schedule_doc(extract("Test1-may4_2026.pdf"), str(tmp_path))
    assert schedule_case_ids(schedule) == ["15419897", "15419898"]


def test_reprocessing_a_case_updates_its_row_instead_of_adding_one(tmp_path):
    create_or_update_schedule_doc(extract("Test1-may4_2026.pdf"), str(tmp_path))
    schedule = create_or_update_schedule_doc(extract("Test1-may4_2026.pdf"), str(tmp_path))
    assert schedule_case_ids(schedule) == ["15419897"]


def test_different_days_get_separate_schedules(tmp_path):
    may4 = create_or_update_schedule_doc(extract("Test1-may4_2026.pdf"), str(tmp_path))
    may10 = create_or_update_schedule_doc(extract("Test4-may10_2026.pdf"), str(tmp_path))
    assert Path(may4) != Path(may10)
    assert schedule_case_ids(may10) == ["15419372"]