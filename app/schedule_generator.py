import os
import re
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH


SCHEDULE_HEADERS = [
    "Case ID",
    "Patient / Doctor / Phone",
    "DOB",
    "Adult",
    "Child",
    "Date & Time",
    "Services Requested",
    "Allegations",
    "Comments",
]

SCHEDULE_WIDTHS = [
    Inches(0.75),  # Case ID
    Inches(1.45),  # Patient / Doctor / Phone
    Inches(0.9),   # DOB
    Inches(0.45),  # Adult
    Inches(0.45),  # Child
    Inches(1.05),  # Date & Time
    Inches(1.85),  # Services Requested
    Inches(1.85),  # Allegations
    Inches(1.55),  # Comments
]


def get_short_doctor_name(full_doctor_name):
    if not full_doctor_name or full_doctor_name == "Unknown Doctor":
        return "Unknown Doctor"

    doctor_name = full_doctor_name.strip()

    if re.fullmatch(r"Dr\.\s*[A-Z]", doctor_name):
        return doctor_name

    cleaned = doctor_name.split(",")[0].strip()
    name_parts = re.findall(r"[A-Za-z]+", cleaned)

    if not name_parts:
        return "Unknown Doctor"

    last_name = name_parts[-1]
    return f"Dr. {last_name[0].upper()}"


def schedule_file_name(schedule_file_date):
    return f"Schedule_{schedule_file_date}.docx"


def build_patient_cell_text(extracted_data):
    patient_name = extracted_data.get("patient_name", "Unknown") or "Unknown"
    doctor_short = get_short_doctor_name(extracted_data.get("doctor_assigned", ""))
    phone = extracted_data.get("phone", "Unknown")

    return (
        f"{patient_name}\n"
        f"{phone}\n"
        f"\n"
        f"{doctor_short}"
    )


def parse_datetime_for_sort(date_time_text):
    """
    Converts:
    May 4th, 2026 08:00 AM MST
    into a sortable datetime object.
    """
    if not date_time_text:
        return datetime.max

    cleaned = re.sub(r"(\d{1,2})(st|nd|rd|th)", r"\1", date_time_text)
    cleaned = cleaned.replace("\n", " ")
    cleaned = cleaned.replace(" MST", "").strip()

    try:
        return datetime.strptime(cleaned, "%B %d, %Y %I:%M %p")
    except ValueError:
        return datetime.max


def set_cell_text(cell, text, bold=False, font_size=9):
    cell.text = ""

    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run = paragraph.add_run(str(text or ""))
    run.bold = bold
    run.font.size = Pt(font_size)


def set_cell_width(cell, width):
    cell.width = width

    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = tcPr.first_child_found_in("w:tcW")

    if tcW is None:
        from docx.oxml import OxmlElement
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)

    tcW.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w", str(int(width.inches * 1440)))
    tcW.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type", "dxa")


def apply_table_formatting(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for row_idx, row in enumerate(table.rows):
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO

        for cell_idx, cell in enumerate(row.cells):
            if cell_idx < len(SCHEDULE_WIDTHS):
                set_cell_width(cell, SCHEDULE_WIDTHS[cell_idx])

            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)

            if row_idx == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.size = Pt(9)


def set_landscape_layout(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE

    section.page_width = Inches(11)
    section.page_height = Inches(8.5)

    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.35)
    section.right_margin = Inches(0.35)


def get_table_rows_as_data(table):
    rows = []

    for row in table.rows[1:]:
        rows.append([cell.text for cell in row.cells])

    return rows


def clear_table_body(table):
    """
    Deletes all rows except the header row.
    """
    while len(table.rows) > 1:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)


def sort_table_by_datetime(table):
    rows = get_table_rows_as_data(table)

    # Date & Time column index is 5
    rows.sort(key=lambda r: parse_datetime_for_sort(r[5] if len(r) > 5 else ""))

    clear_table_body(table)

    for row_data in rows:
        row = table.add_row().cells

        for i, value in enumerate(row_data):
            if i < len(row):
                row[i].text = value


def ensure_schedule_table_columns(table):
    """
    Adds newer schedule columns to older generated DOCX files.
    """
    while len(table.columns) < len(SCHEDULE_HEADERS):
        column_index = len(table.columns)
        width = SCHEDULE_WIDTHS[column_index]
        table.add_column(width)

    for i, header in enumerate(SCHEDULE_HEADERS):
        set_cell_text(table.rows[0].cells[i], header, bold=True, font_size=9)


def build_date_time_text(extracted_data):
    return (
        f"{extracted_data.get('schedule_date', '')}\n "
        f"{extracted_data.get('schedule_time', '')}"
    ).strip()


def populate_schedule_row(row, extracted_data):
    set_cell_text(row[0], extracted_data.get("case_id", ""))
    set_cell_text(row[1], build_patient_cell_text(extracted_data))
    set_cell_text(row[2], extracted_data.get("dob", ""))
    set_cell_text(row[3], extracted_data.get("adult", ""))
    set_cell_text(row[4], extracted_data.get("child", ""))
    set_cell_text(row[5], build_date_time_text(extracted_data))
    set_cell_text(row[6], extracted_data.get("services_requested", ""))
    set_cell_text(row[7], extracted_data.get("allegations", "Not Available"))
    set_cell_text(row[8], "")


def create_or_update_schedule_doc(extracted_data, schedules_folder):
    os.makedirs(schedules_folder, exist_ok=True)

    file_name = schedule_file_name(extracted_data["schedule_file_date"])
    file_path = os.path.join(schedules_folder, file_name)

    if os.path.exists(file_path):
        doc = Document(file_path)
        set_landscape_layout(doc)
    else:
        doc = Document()
        set_landscape_layout(doc)

        doc.add_heading(f"Schedule for {extracted_data['schedule_date']}", 0)

        table = doc.add_table(rows=1, cols=len(SCHEDULE_HEADERS))
        table.style = "Table Grid"

        for i, header in enumerate(SCHEDULE_HEADERS):
            set_cell_text(table.rows[0].cells[i], header, bold=True, font_size=9)

    table = doc.tables[0]
    ensure_schedule_table_columns(table)

    case_id = extracted_data.get("case_id", "")

    for existing_row in table.rows[1:]:
        if existing_row.cells[0].text.strip() == case_id:
            print(f"Updating existing case ID in schedule: {case_id}")
            populate_schedule_row(existing_row.cells, extracted_data)
            apply_table_formatting(table)
            sort_table_by_datetime(table)
            apply_table_formatting(table)
            doc.save(file_path)
            return file_path

    row = table.add_row().cells
    populate_schedule_row(row, extracted_data)

    sort_table_by_datetime(table)
    apply_table_formatting(table)

    doc.save(file_path)

    print(f"Schedule created/updated: {file_path}")
    return file_path
