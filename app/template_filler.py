import os
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from datetime import datetime



FONT_NAME = "Garamond"
FONT_SIZE = 10

RED_PLACEHOLDERS = {
    "{{EXAMINER}}",
    "{{DOCTOR_FULL_NAME}}"
}


def clean_filename(text):
    text = str(text or "").strip()
    text = re.sub(r"[^\w\s-]", "", text)
    return text.replace(" ", "_")


def normalize_name_part(name_part):
    pieces = re.split(r"([-'’])", name_part.lower())
    return "".join(
        piece.capitalize() if piece not in {"-", "'", "’"} else piece
        for piece in pieces
    )


def normalize_patient_name(full_name):
    if not full_name or full_name.strip().lower() == "unknown":
        return "Unknown"

    parts = full_name.strip().split()
    suffixes = {"jr", "sr", "ii", "iii", "iv"}

    return " ".join(
        part.upper() if part.lower().strip(".") in suffixes else normalize_name_part(part)
        for part in parts
    )


def get_first_name(full_name):
    if not full_name or full_name.strip().lower() == "unknown":
        return "Unknown"

    parts = full_name.strip().split()
    return parts[0] if parts else "Unknown"

def get_last_name(full_name):
    if not full_name or full_name.strip().lower() == "unknown":
        return "Unknown"

    parts = full_name.strip().split()

    if len(parts) == 1:
        return parts[0]

    return parts[-1]

def clean_doctor_full_name(full_doctor_name):
    """
    JONATHAN SHELTON, PSYD -> JONATHAN SHELTON
    """
    if not full_doctor_name or full_doctor_name == "Unknown Doctor":
        return "Unknown Doctor"

    return full_doctor_name.split(",")[0].strip()


def format_report_filename(full_name):
    """
    John Doe -> Doe_John_report.docx
    Unknown -> Unknown_Patient_report.docx
    """
    if not full_name or full_name.strip().lower() == "unknown":
        return "Unknown_Patient_report.docx"

    parts = full_name.strip().split()

    if len(parts) == 1:
        return f"{clean_filename(parts[0])}_report.docx"

    first_name = parts[0]
    last_name = parts[-1]

    return f"{clean_filename(last_name)}_{clean_filename(first_name)}_report.docx"


def apply_font_to_run(run, red=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(FONT_SIZE)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)

    # Only force red for these placeholders.
    # Otherwise, keep whatever color already exists in the Word template.
    if red:
        run.font.color.rgb = RGBColor(255, 0, 0)


def replace_placeholder_across_runs(paragraph, placeholder, value, red=False):
    """
    Replaces a placeholder even when Word has split it across multiple runs.
    The replacement is placed in the first run touched by the placeholder so
    surrounding formatting is preserved as much as Word allows.
    """
    replacement = str(value or "")

    while True:
        run_texts = [run.text for run in paragraph.runs]
        full_text = "".join(run_texts)
        start = full_text.find(placeholder)

        if start == -1:
            return False

        end = start + len(placeholder)
        current_pos = 0
        start_run_idx = None
        start_offset = None
        end_run_idx = None
        end_offset = None

        for idx, text in enumerate(run_texts):
            next_pos = current_pos + len(text)

            if start_run_idx is None and start < next_pos:
                start_run_idx = idx
                start_offset = start - current_pos

            if end_run_idx is None and end <= next_pos:
                end_run_idx = idx
                end_offset = end - current_pos
                break

            current_pos = next_pos

        if start_run_idx is None or end_run_idx is None:
            return False

        start_run = paragraph.runs[start_run_idx]
        end_run = paragraph.runs[end_run_idx]

        prefix = start_run.text[:start_offset]
        suffix = end_run.text[end_offset:]

        if start_run_idx == end_run_idx:
            start_run.text = f"{prefix}{replacement}{suffix}"
        else:
            start_run.text = f"{prefix}{replacement}"

            for idx in range(start_run_idx + 1, end_run_idx):
                paragraph.runs[idx].text = ""

            end_run.text = suffix

        apply_font_to_run(start_run, red=red)


def replace_placeholders_in_paragraph(paragraph, replacements):
    """
    Replaces placeholders inside existing runs so manual formatting
    in the Word template is preserved.

    This prevents red template text from turning black.
    """

    for placeholder, value in replacements.items():
        replace_placeholder_across_runs(
            paragraph,
            placeholder,
            value,
            red=placeholder in RED_PLACEHOLDERS
        )

def replace_placeholders_in_table(table, replacements):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_placeholders_in_paragraph(paragraph, replacements)


def replace_placeholders_in_headers_and_footers(doc, replacements):
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            replace_placeholders_in_paragraph(paragraph, replacements)

        for table in section.header.tables:
            replace_placeholders_in_table(table, replacements)

        for paragraph in section.footer.paragraphs:
            replace_placeholders_in_paragraph(paragraph, replacements)

        for table in section.footer.tables:
            replace_placeholders_in_table(table, replacements)

def clean_ordinal_date(date_text):
    """
    Converts May 10th, 2026 -> May 10, 2026
    """
    return re.sub(r"(\d{1,2})(st|nd|rd|th)", r"\1", str(date_text or ""))


def parse_month_day_year(date_text):
    """
    Parses:
    November 4, 2011
    May 10th, 2026
    """
    cleaned = clean_ordinal_date(date_text).strip()

    try:
        return datetime.strptime(cleaned, "%B %d, %Y")
    except ValueError:
        return None


def calculate_years_months(dob_dt, exam_dt):
    """
    Calculates age as years and months.
    """
    years = exam_dt.year - dob_dt.year
    months = exam_dt.month - dob_dt.month

    if exam_dt.day < dob_dt.day:
        months -= 1

    if months < 0:
        years -= 1
        months += 12

    return years, months


def format_dob_with_age(dob_text, exam_date_text):
    """
    Converts:
    November 4, 2011 + May 10th, 2026
    into:
    11/04/2011; 14 years, 6 months
    """
    dob_dt = parse_month_day_year(dob_text)
    exam_dt = parse_month_day_year(exam_date_text)

    if not dob_dt:
        return dob_text or ""

    formatted_dob = dob_dt.strftime("%m/%d/%Y")

    if not exam_dt:
        return formatted_dob

    years, months = calculate_years_months(dob_dt, exam_dt)

    year_label = "year" if years == 1 else "years"
    month_label = "month" if months == 1 else "months"

    return f"{formatted_dob}; {years} {year_label}, {months} {month_label}"

def format_exam_date_mmddyyyy(exam_date_text):
    """
    Converts:
    May 10th, 2026 -> 05/10/2026
    """
    exam_dt = parse_month_day_year(exam_date_text)

    if not exam_dt:
        return exam_date_text or ""

    return exam_dt.strftime("%m/%d/%Y")

def fill_report_template(extracted_data, template_path, output_folder):
    os.makedirs(output_folder, exist_ok=True)

    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template not found: {template_path}")

    full_name = normalize_patient_name(
        extracted_data.get("patient_name", "Unknown") or "Unknown"
    )
    first_name = get_first_name(full_name)
    last_name = get_last_name(full_name)

    exam_date_formatted = format_exam_date_mmddyyyy(
        extracted_data.get("schedule_date", "")
    )
    dob_formatted = format_dob_with_age(
        extracted_data.get("dob", ""),
        extracted_data.get("schedule_date", "")
    )

    doctor_full_name = clean_doctor_full_name(
        extracted_data.get("doctor_assigned", "")
    )

    referred_by = extracted_data.get(
        "referred_by",
        "Arizona Department of Economic Security - Disability Determination Service"
    )

    replacements = {
        "{{FULL_NAME}}": full_name,
        "{{FIRST_NAME}}": f"{first_name}",
        "{{First_name}}": f"{first_name}",
        "{{First_Name}}": f"{first_name}",
        "{{first_name}}": f"{first_name}",
        "{{LAST_NAME}}": last_name,
        "{{CASE_NUMBER}}": extracted_data.get("case_id", ""),
        "{{DOB}}": dob_formatted,
        "{{DATE_OF_EXAMINATION}}": exam_date_formatted,
        "{{REFERRED_BY}}": referred_by,
        "{{ALLEGATIONS}}": extracted_data.get("allegations", ""),
        "{{EXAMINER}}": extracted_data.get("examiner", ""),
        "{{DOCTOR_FULL_NAME}}": doctor_full_name
    }

    doc = Document(template_path)

    for paragraph in doc.paragraphs:
        replace_placeholders_in_paragraph(paragraph, replacements)

    for table in doc.tables:
        replace_placeholders_in_table(table, replacements)

    replace_placeholders_in_headers_and_footers(doc, replacements)

    output_filename = format_report_filename(full_name)
    output_path = os.path.join(output_folder, output_filename)

    doc.save(output_path)

    print(f"Generated local report: {output_path}")
    return output_path
